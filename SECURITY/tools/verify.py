# -*- coding: utf-8 -*-
"""v2 산출물 검사: python tools\\verify.py 07  (여러 장: 01 02 ...)"""
import sys, glob, os, re
import docx, pptx

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BAN_ALL = ["**", "실측", "가이드 결함", "우리 환경", "여러분"]
BAN_PPT = ["`"]
TIC = ["것이 핵심이다", "라는 뜻이다", "라는 뜻이다."]
NO_NOTE = ("학습 목표", "Agenda", "이 장의 구성", "요약", "퀴즈", "Quiz", "실습 개요")


def check(nn):
    d = os.path.join(ROOT, nn + "장")
    bad = []
    dx = [f for f in glob.glob(os.path.join(d, "*.docx"))]
    px = [f for f in glob.glob(os.path.join(d, "*.pptx"))]
    if len(dx) != 1 or len(px) != 1:
        return ["docx/pptx 가 각각 1개가 아님: %d/%d" % (len(dx), len(px))]
    doc = docx.Document(dx[0])
    paras = [p.text for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    n = sum(map(len, paras)) + sum(map(len, cells))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    text = "\n".join(paras + cells)
    if not 26000 <= n <= 30000:
        bad.append("docx 글자 수 %d" % n)
    if len(h1) != 1:
        bad.append("H1 %d개" % len(h1))
    for k in ("학습 목표", "요약", "퀴즈", "실습 개요"):
        if not any(k in h for h in h2):
            bad.append("H2 '%s' 없음" % k)
    for w in BAN_ALL:
        if w in text:
            bad.append("docx 금지어 %r %d건" % (w, text.count(w)))
    tic = sum(text.count(t) for t in TIC)
    if tic > 2:
        bad.append("docx 재진술 어미 %d건" % tic)
    prs = pptx.Presentation(px[0])
    ns = len(prs.slides)
    if not 35 <= ns <= 55:
        bad.append("슬라이드 %d장" % ns)
    allt = []
    for i, s in enumerate(prs.slides, 1):
        ts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        allt += ts
        note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        allt.append(note)
        title = ts[0] if ts else ""
        need = i > 1 and not any(k in " ".join(ts[:1]) for k in NO_NOTE) and not is_section(s)
        if need and not 200 <= len(note) <= 500:
            bad.append("슬라이드 %d 노트 %d자 (%s)" % (i, len(note), title[:20]))
    pt = "\n".join(allt)
    for w in BAN_ALL + BAN_PPT:
        if w in pt:
            bad.append("pptx 금지어 %r %d건" % (w, pt.count(w)))
    tx = sorted(glob.glob(os.path.join(d, "*_실습_*.txt")))
    if not 3 <= len(tx) <= 8:
        bad.append("트랜스크립트 %d개" % len(tx))
    for f in tx:
        t = open(f, encoding="utf-8").read()
        if not re.search(r"(ORA|TNS|PLS|SP2|PKI)-\d{4,5}", t):
            bad.append("오류 흐름 없음: %s" % os.path.basename(f))
        if not os.path.basename(f).startswith(str(int(nn)) + "장_"):
            bad.append("장 번호 불일치: %s" % os.path.basename(f))
    print("%s장  트랜스크립트 %d개 /" % (nn, len(tx)), end=" ")
    print("%s장  docx %d자 H2 %d개 / pptx %d장" % (nn, n, len(h2), ns))
    return bad


def is_section(s):
    # 섹션 슬라이드: 전면 붉은 사각형 + 번호/제목/선/불릿
    sh = list(s.shapes)
    return len(sh) == 5 and sh[0].width == 12191695 and sh[0].height == 6858000 \
        and str(sh[0].fill.fore_color.rgb) == "C0392B"


if __name__ == "__main__":
    for nn in sys.argv[1:]:
        r = check(nn.zfill(2))
        print("  OK" if not r else "\n".join("  ✗ " + x for x in r))
