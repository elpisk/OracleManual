# -*- coding: utf-8 -*-
"""5일 과정 교안 검사:  python3 tools/verify.py 01 02 ...   (인자 없으면 전체)"""
import sys, glob, os, re
import docx, pptx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BAN_ALL = ["**", "실측", "우리 환경", "여러분", "Day ", "Day1", "Day2", "Day3", "Day4", "Day5",
           "일차", "첫째 날", "둘째 날", "셋째 날", "넷째 날", "다섯째 날", "오전 수업", "오후 수업"]
BAN_PPT = ["`"]
OLD_ENV = ["19.0.0/grid", "dbhome_1", "oracleasm", "ORCL:", "oel7rac", "OEL7V9", "+ASM3",
           "racstby", "racsby"]
INSTALL = ["gridSetup.sh -silent", "runInstaller", "-createDatabase", "runcluvfy",
           "stage -pre crsinst", "udevadm"]
TIC = ["것이 핵심이다", "라는 뜻이다"]
NO_NOTE = ("학습 목표", "Agenda", "요약", "퀴즈", "Quiz", "실습 개요")
DOCX_RANGE = (14000, 22000)
PPT_RANGE = (28, 45)


def is_section(s):
    sh = list(s.shapes)
    return len(sh) == 5 and sh[0].width == 12191695 and sh[0].height == 6858000 \
        and str(sh[0].fill.fore_color.rgb) == "C0392B"


def check(nn):
    d = os.path.join(ROOT, nn + "장")
    bad = []
    dx = glob.glob(os.path.join(d, "*.docx"))
    px = glob.glob(os.path.join(d, "*.pptx"))
    if len(dx) != 1 or len(px) != 1:
        return ["docx/pptx 가 각각 1개가 아님: %d/%d" % (len(dx), len(px))]
    for f in dx + px:
        if not os.path.basename(f).startswith("%d장_" % int(nn)):
            bad.append("파일 장 번호 불일치: " + os.path.basename(f))
    doc = docx.Document(dx[0])
    paras = [p.text for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    n = sum(map(len, paras)) + sum(map(len, cells))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    text = re.sub(r"RAC/\S+\.txt|\d+장_실습_\S+", "", "\n".join(paras + cells))
    nimg = len(doc.inline_shapes)
    if not DOCX_RANGE[0] <= n <= DOCX_RANGE[1]:
        bad.append("docx 글자 수 %d (기준 %d~%d)" % ((n,) + DOCX_RANGE))
    if len(h1) != 1 or not h1[0].startswith("제%d장" % int(nn)):
        bad.append("H1 이상: %r" % h1)
    for h in h2:
        if not h.startswith("%d." % int(nn)):
            bad.append("H2 번호 불일치: " + h[:30])
    for k in ("학습 목표", "요약", "퀴즈", "실습 개요"):
        if not any(k in h for h in h2):
            bad.append("H2 '%s' 없음" % k)
    if nimg < 4:
        bad.append("docx 그림 %d개 (4개 이상)" % nimg)
    for w in BAN_ALL + OLD_ENV + INSTALL:
        if w in text:
            bad.append("docx 금지어 %r %d건" % (w, text.count(w)))
    tic = sum(text.count(t) for t in TIC)
    if tic > 2:
        bad.append("docx 재진술 어미 %d건" % tic)
    prs = pptx.Presentation(px[0])
    ns = len(prs.slides)
    if not PPT_RANGE[0] <= ns <= PPT_RANGE[1]:
        bad.append("슬라이드 %d장 (기준 %d~%d)" % ((ns,) + PPT_RANGE))
    allt, npic = [], 0
    for i, s in enumerate(prs.slides, 1):
        ts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        npic += sum(1 for sh in s.shapes if sh.shape_type == 13)
        allt += ts
        note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        allt.append(note)
        title = ts[0] if ts else ""
        need = i > 1 and not any(k in " ".join(ts[:1]) for k in NO_NOTE) and not is_section(s)
        if need and not 200 <= len(note) <= 500:
            bad.append("슬라이드 %d 노트 %d자 (%s)" % (i, len(note), title[:20]))
    if i > 0 and str(prs.slides[0].shapes[3].text_frame.text) != "Oracle Database 19c RAC 과정":
        bad.append("PPT 과정명 불일치")
    if npic < 4:
        bad.append("pptx 그림 %d개 (4개 이상)" % npic)
    pt = re.sub(r"RAC/\S+\.txt|\d+장_실습_\S+", "", "\n".join(allt))
    for w in BAN_ALL + BAN_PPT + OLD_ENV + INSTALL:
        if w in pt:
            bad.append("pptx 금지어 %r %d건" % (w, pt.count(w)))
    print("%s장  docx %d자 H2 %d개 그림 %d / pptx %d장 그림 %d" % (nn, n, len(h2), nimg, ns, npic))
    return bad


if __name__ == "__main__":
    args = sys.argv[1:] or sorted(os.path.basename(x)[:2] for x in glob.glob(os.path.join(ROOT, "[0-9][0-9]장")))
    fail = 0
    for nn in args:
        r = check(nn.zfill(2))
        fail += bool(r)
        print("  OK" if not r else "\n".join("  ✗ " + x for x in r))
    sys.exit(1 if fail else 0)
