# -*- coding: utf-8 -*-
"""교안 docx 빌더 — 기존 챕터 docx 를 템플릿으로 열어 서식(스타일 정의)을 그대로 물려받고,
본문만 비운 뒤 마크업 텍스트로 새로 채운다.

마크업 태그 (줄 맨 앞)
  @H1 / @H2 / @H3   제목
  @P                본문 문단
  @B                불릿(List Bullet)
  @CODE ... @END    검은 코드 블록 (Consolas 9.5, 배경 1E1E1E)
  @NOTE 제목        노랑 박스   ... @END
  @BP   제목        녹색 박스   ... @END
  @CASE 제목        파랑 박스   ... @END
  @QUIZ 제목        보라 박스   ... @END
  @TBL ... @END     표(첫 줄이 머리글, 칸 구분자 |)
  @BOLD             굵은 문단 (정답지의 문항 줄)
"""
import io, re, copy, docx
from docx.shared import Pt, RGBColor, Emu, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CODE_BG = "1E1E1E"
CODE_FG = RGBColor(0xF2, 0xF2, 0xF2)
BOX = {"NOTE": ("FFF6D5", "Note"), "BP": ("E3F0DB", "Best Practice"),
       "CASE": ("E1ECF7", "실무 시나리오"), "QUIZ": ("F3E4F5", "Quiz")}
TBL_HDR = "D9E1F2"
IMGDIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "img")
BODY_W = Emu(5486400)

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')


def _runs(p, text, size, bold=False):
    """**굵게** 는 굵은 글씨, `코드` 는 Consolas 로 바꾸고 표시 기호는 남기지 않는다."""
    for part in INLINE.split(text):
        if not part:
            continue
        b, mono = bold, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        r = p.add_run(part)
        fn = "Consolas" if mono else "맑은 고딕"
        r.font.name = fn
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), "맑은 고딕")
        r.font.size = Pt(size)
        if b:
            r.bold = True
    return p


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    tcPr.append(s)


class Doc(object):
    def __init__(self, template):
        self.d = docx.Document(template)
        body = self.d.element.body
        for child in list(body):
            if child.tag == qn('w:sectPr'):
                continue
            body.remove(child)

    # ── 기본 요소 ────────────────────────────────────────────
    def head(self, level, text):
        self.d.add_paragraph(text, style="Heading %d" % level)

    def para(self, text, bold=False):
        p = self.d.add_paragraph()
        return _runs(p, text, 10.5, bold)

    def bullet(self, text):
        p = self.d.add_paragraph(style="List Bullet")
        return _runs(p, text, 10.5)

    def code(self, lines):
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.columns[0].width = BODY_W
        c = t.rows[0].cells[0]
        c.width = BODY_W
        _shade(c, CODE_BG)
        first = True
        for ln in lines:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(ln if ln.strip() else " ")
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn('w:eastAsia'), "Consolas")
            r.font.size = Pt(9.5)
            r.font.color.rgb = CODE_FG
        self.d.add_paragraph()

    def box(self, kind, title, lines):
        fill, label = BOX[kind]
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.columns[0].width = BODY_W
        c = t.rows[0].cells[0]
        c.width = BODY_W
        _shade(c, fill)
        p = c.paragraphs[0]
        r = p.add_run("%s · %s" % (label, title))
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), "맑은 고딕")
        r.font.size = Pt(10.5)
        r.bold = True
        for ln in lines:
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _runs(p, ln, 10)
        self.d.add_paragraph()

    def image(self, path, caption, width=6.3):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Inches(width))
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(12)
        r = c.add_run(caption)
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), "맑은 고딕")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        r.italic = True

    def table(self, rows):
        ncol = len(rows[0])
        t = self.d.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.rows[i].cells[j]
                if i == 0:
                    _shade(c, TBL_HDR)
                p = c.paragraphs[0]
                _runs(p, val, 9.5, bold=(i == 0))
        self.d.add_paragraph()

    def save(self, path):
        self.d.save(path)


# ── 마크업 처리 ──────────────────────────────────────────────
def build(markup_path, template, out_path):
    src = io.open(markup_path, encoding="utf-8").read().split("\n")
    doc = Doc(template)
    i, n = 0, len(src)
    while i < n:
        ln = src[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        m = re.match(r'^@(H1|H2|H3|P|BOLD|B)(?:\s+(.*))?$', s)
        if m:
            tag, txt = m.group(1), m.group(2)
            if tag == "H1":
                doc.head(1, txt)
            elif tag == "H2":
                doc.head(2, txt)
            elif tag == "H3":
                doc.head(3, txt)
            elif tag == "P":
                doc.para(txt)
            elif tag == "BOLD":
                doc.para(txt, bold=True)
            else:
                doc.bullet(txt)
            i += 1
            continue
        m = re.match(r'^@IMG\s+(\S+)\s*\|\s*(.*)$', s)
        if m:
            doc.image(os.path.join(IMGDIR, m.group(1)), m.group(2))
            i += 1
            continue
        m = re.match(r'^@(CODE|TBL)\s*$', s)
        if m:
            tag = m.group(1)
            buf = []
            i += 1
            while i < n and src[i].strip() != "@END":
                buf.append(src[i].rstrip("\n"))
                i += 1
            i += 1
            if tag == "CODE":
                doc.code(buf)
            else:
                doc.table([[c.strip() for c in r.split("|")] for r in buf if r.strip()])
            continue
        m = re.match(r'^@(NOTE|BP|CASE|QUIZ)\s+(.*)$', s)
        if m:
            kind, title = m.group(1), m.group(2)
            buf = []
            i += 1
            while i < n and src[i].strip() != "@END":
                t = src[i].strip()
                if t:
                    if t.startswith("@P "):
                        t = t[3:]
                    buf.append(t)
                i += 1
            i += 1
            doc.box(kind, title, buf)
            continue
        raise ValueError("알 수 없는 줄 %d: %r" % (i + 1, s[:60]))
    doc.save(out_path)
    d = docx.Document(out_path)
    tot = sum(len(p.text) for p in d.paragraphs)
    tot += sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells)
    return tot, len(d.paragraphs), len(d.tables)


if __name__ == "__main__":
    import sys
    tot, np_, nt = build(sys.argv[1], sys.argv[2], sys.argv[3])
    print("%s  →  %d자 / 문단 %d / 표 %d" % (sys.argv[3], tot, np_, nt))
